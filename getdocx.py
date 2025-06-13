from docx import Document
from openpyxl import load_workbook

wb = load_workbook("openpyxl_output.xlsx")
sheet = wb["shuju"]
text_list = []
for row in sheet['A2':'H41']:
    text = ""
    for cell in row:
        text += str(cell.value)
    #print(text)
    char1 = "专利"
    char2 = "."
    char3 = "（详见第XX页）"
    s = char1 + text
    t = s[:4] + char2 + s[4:]
    new_text = t[:len(t)]+char3+t[len(t):]
    text_list.append(new_text)

doc = Document()

for text in text_list:
    doc.add_paragraph(text)
doc.save("氮化镓专利成果41-80.docx")


